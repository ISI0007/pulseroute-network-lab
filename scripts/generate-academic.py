"""Generate evidence-aware academic artifacts from versioned project content."""
import json, math, pathlib, textwrap
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch, FancyArrowPatch
from docx import Document
from docx.shared import Inches as DI, Pt as DP, RGBColor as DC
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak, Image, Table, TableStyle, KeepTogether
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from xml.sax.saxutils import escape

ROOT=pathlib.Path(__file__).resolve().parents[1]
OUT=ROOT/"deliverables";OUT.mkdir(exist_ok=True)
FIG=ROOT/"docs/diagrams";FIG.mkdir(exist_ok=True)
data=json.loads((ROOT/"docs/academic-content.json").read_text(encoding="utf-8"))
BG="#0b1019"; PANEL="#152131"; CYAN="#60e2d2"; TEXT="#edf3f8"; MUTED="#a7b6c9"
plt.rcParams.update({"font.family":"DejaVu Sans","svg.fonttype":"none"})

def graph(d):
    fig,ax=plt.subplots(figsize=(12,7),facecolor=BG);ax.set_facecolor(BG);ax.set_xlim(-.4,12.4);ax.set_ylim(-.4,7.2);ax.axis("off")
    ax.text(0,6.9,d["title"],color=TEXT,fontsize=19,weight="bold")
    ax.text(0,6.5,"PULSEROUTE / IMPLEMENTATION MODEL",color=CYAN,fontsize=9)
    n=len(d["nodes"])
    layouts={
     "01":[(1.2,4.8),(4.3,4.8),(7.4,4.8),(7.4,1.6),(4.3,1.6),(10.7,3.2)],
     "02":[(1.2,3.4),(5,3.4),(10.4,5),(10.4,1.5),(5,1)],
     "03":[(1.2,5),(5.6,5),(10.4,5),(10.4,1.5),(5.6,1.5),(1.2,1.5)],
     "04":[(1.7,5),(6,5),(10.3,5),(3.4,1.8),(8.3,1.8)],
     "05":[(1.7,3.3),(8.8,5.4),(8.8,4.3),(8.8,3.2),(8.8,2.1),(8.8,1)],
     "07":[(1.2,5.2),(5.2,5.2),(9.7,5.2),(9.7,3.1),(5.2,3.1),(1.2,3.1),(9.7,.8),(5.2,.8)],
     "08":[(1.4,4.9),(6,4.9),(10.6,4.9),(4.2,1.5),(9,1.5)],
     "09":[(1.2,5.2),(5.2,5.2),(9.7,5.2),(9.7,3.1),(5.2,3.1),(1.2,3.1),(1.2,.8),(7.3,.8)],
     "10":[(1.2,4.8),(5,4.8),(5,1.6),(9.8,1.6),(9.8,4.8),(10.6,3.2)]
    }
    positions=layouts[d["id"][:2]]
    for a,b,label in d["edges"]:
        x1,y1=positions[a];x2,y2=positions[b]
        dx,dy=x2-x1,y2-y1;dist=math.hypot(dx,dy)
        scale=max(abs(dx)/1.29,abs(dy)/.65)
        start=(x1+dx/scale,y1+dy/scale);end=(x2-dx/scale,y2-dy/scale)
        arrow=FancyArrowPatch(start,end,arrowstyle="-|>",mutation_scale=13,color=CYAN,lw=1.4,connectionstyle="arc3,rad=0.02")
        ax.add_patch(arrow)
        ax.text((start[0]+end[0])/2,(start[1]+end[1])/2+(.78 if abs(dy)<.2 else .16),textwrap.fill(label,20),ha="center",va="center",fontsize=8,color=MUTED,bbox={"facecolor":BG,"edgecolor":"none","pad":2})
    for i,(x,y) in enumerate(positions):
        label=d["nodes"][i].replace("\\n","\n")
        lines="\n".join(textwrap.fill(line,19) for line in label.splitlines())
        box=FancyBboxPatch((x-1.15,y-.48),2.3,.96,boxstyle="round,pad=.08,rounding_size=.10",edgecolor="#34465e",facecolor=PANEL,lw=1.1)
        ax.add_patch(box);ax.text(x,y,lines,ha="center",va="center",color=TEXT,fontsize=9.5)
    ax.text(0,-.15,"Design/source diagram. Runtime verification status is documented separately.",color=MUTED,fontsize=8)
    for ext in ("png","svg"):fig.savefig(FIG/(d["id"]+"."+ext),dpi=180,bbox_inches="tight",facecolor=BG)
    plt.close(fig)

def sequence(d):
    fig,ax=plt.subplots(figsize=(12,8),facecolor=BG);ax.set_xlim(0,10);ax.set_ylim(0,13);ax.axis("off")
    ax.text(.3,12.5,d["title"],color=TEXT,fontsize=19,weight="bold")
    names=["Sender","Receiver","Disk","SQLite"];xs=[1,3.7,6.4,9]
    for x,name in zip(xs,names):
        ax.add_patch(FancyBboxPatch((x-.7,11.2),1.4,.65,boxstyle="round,pad=.1",facecolor=PANEL,edgecolor=CYAN))
        ax.text(x,11.52,name,ha="center",va="center",color=TEXT,fontsize=12)
        ax.plot([x,x],[1,11.1],color="#34465e",linestyle="--")
    messages=[(0,0,"Hash source, rewind"),(0,1,"put(id, name, size, SHA-256)"),(1,0,"ready"),(0,1,"binary chunk (repeat)"),(1,2,"write temporary + update hash"),(0,1,"end"),(1,1,"verify length and SHA-256"),(1,2,"sync / close / unique rename"),(1,3,"record verified transfer"),(1,0,"complete(id, size, SHA-256)")]
    for row,(a,b,label) in enumerate(messages):
        y=10.5-row*.94
        if a==b:
            ax.text(xs[a]+.15,y,label,color=CYAN,fontsize=10,ha="left")
        else:
            ax.annotate("",xy=(xs[b],y),xytext=(xs[a],y),arrowprops={"arrowstyle":"->","color":CYAN,"lw":1.5})
            ax.text((xs[a]+xs[b])/2,y+.18,label,ha="center",color=TEXT,fontsize=10,bbox={"facecolor":BG,"edgecolor":"none","pad":2})
    ax.text(.3,.35,"Failure paths remove temporary data; SQL failure after rename removes the completed file.\nA crash between filesystem and database operations still requires future reconciliation.",color=MUTED,fontsize=9)
    for ext in ("png","svg"):fig.savefig(FIG/(d["id"]+"."+ext),dpi=180,bbox_inches="tight",facecolor=BG)
    plt.close(fig)

for d in data["diagrams"]:
    (FIG/(d["id"]+".mmd")).write_text(d["mermaid"],encoding="utf-8")
    sequence(d) if d["id"].startswith("06") else graph(d)

# A neutral diagram figure is not a fabricated benchmark chart.
doc=Document()
for sec in doc.sections:
    sec.top_margin=DI(.75);sec.bottom_margin=DI(.7)
normal=doc.styles["Normal"];normal.font.name="Calibri";normal.font.size=DP(11);normal.paragraph_format.space_after=DP(8)
for level in ("Title","Heading 1","Heading 2"):doc.styles[level].font.color.rgb=DC.from_string("126F68")
doc.add_heading("PULSEROUTE",0);doc.add_paragraph(data["subtitle"],style="Subtitle")
doc.add_paragraph(data["status"]);doc.add_paragraph(data["date"]);doc.add_paragraph("Project owner: ISI0007\nUniversity / course / student identity: to be supplied by the author.")
doc.add_paragraph("Verification statement: historical core tests and current permitted checks are reported separately. No performance measurements or native desktop success are invented.")
doc.add_page_break();doc.add_heading("Contents",1)
for i,(title,_) in enumerate(data["sections"],1):doc.add_paragraph(f"{i}. {title}")
doc.add_page_break()
fig_by_section={8:"01-system-architecture",9:"06-transfer-sequence",11:"04-entity-relationship",16:"08-network-topology"}
for i,(title,body) in enumerate(data["sections"],1):
    doc.add_heading(f"{i}. {title}",1)
    for paragraph in body.replace("\\n","\n").split("\n"):doc.add_paragraph(paragraph)
    if i in fig_by_section:
        doc.add_picture(str(FIG/(fig_by_section[i]+".png")),width=DI(6.2))
        doc.add_paragraph("Source-derived diagram; see editable Mermaid and SVG files.",style="Caption")
doc.add_page_break();doc.add_heading("Frontend evidence",1)
doc.add_paragraph("These are genuine production-frontend screenshots in Chrome. The Go desktop bridge was unavailable; the warning is intentionally visible. They do not prove native networking workflows.")
shot=ROOT/"docs/evidence/frontend-overview.png"
if shot.exists():doc.add_picture(str(shot),width=DI(6.1))
doc.save(OUT/"PulseRoute-Technical-Report-DRAFT.docx")

styles=getSampleStyleSheet()
styles.add(ParagraphStyle(name="Cover",fontName="Helvetica-Bold",fontSize=32,leading=38,textColor=colors.HexColor("#126f68"),spaceAfter=22))
styles["BodyText"].fontSize=10.5;styles["BodyText"].leading=15;styles["BodyText"].spaceAfter=10
styles["Heading1"].textColor=colors.HexColor("#126f68");styles["Heading1"].spaceBefore=15;styles["Heading1"].keepWithNext=True
pdf=SimpleDocTemplate(str(OUT/"PulseRoute-Technical-Report-DRAFT.pdf"),pagesize=(595,842),leftMargin=48,rightMargin=48,topMargin=48,bottomMargin=48)
story=[Spacer(1,65),Paragraph("PULSEROUTE",styles["Cover"]),Paragraph(escape(data["subtitle"]),styles["Title"]),Spacer(1,25),Paragraph(escape(data["status"]),styles["Heading2"]),Paragraph(data["date"],styles["BodyText"]),Paragraph("Project owner: ISI0007",styles["BodyText"]),Paragraph("University, course and student identity are to be supplied by the author.",styles["BodyText"]),Spacer(1,25),Paragraph("Evidence boundary: this is a progress report. Current Go runtime validation is blocked by Smart App Control. Compilation, frontend rendering and schema checks are distinct from full-stack verification.",styles["BodyText"]),PageBreak(),Paragraph("Contents",styles["Heading1"])]
for i,(title,_) in enumerate(data["sections"],1):story.append(Paragraph(f"{i}. {escape(title)}",styles["BodyText"]))
story.append(PageBreak())
for i,(title,body) in enumerate(data["sections"],1):
    story.append(Paragraph(f"{i}. {escape(title)}",styles["Heading1"]))
    for p in body.replace("\\n","\n").split("\n"):story.append(Paragraph(escape(p),styles["BodyText"]))
    if i in fig_by_section:
        img=Image(str(FIG/(fig_by_section[i]+".png")));img.drawHeight=img.imageHeight/img.imageWidth*490;img.drawWidth=490
        story.extend([img,Spacer(1,12)])
story.extend([PageBreak(),Paragraph("Frontend evidence",styles["Heading1"]),Paragraph("Production frontend in Chrome; no mocked data; native Go bridge unavailable.",styles["BodyText"])])
if shot.exists():
    img=Image(str(shot));iw,ih=img.imageWidth,img.imageHeight;img.drawWidth=480;img.drawHeight=ih/iw*480;story.append(img)
def footer(canvas,document):
    canvas.setFont("Helvetica",8);canvas.setFillColor(colors.HexColor("#596675"))
    canvas.drawString(48,25,"PulseRoute | Engineering progress report | 28 October 2026")
    canvas.drawRightString(547,25,str(document.page))
pdf.build(story,onFirstPage=footer,onLaterPages=footer)

# Editable, restrained 16:9 presentation with notes on every slide.
prs=Presentation();prs.slide_width=Inches(13.333);prs.slide_height=Inches(7.5)
def rgb(h):return RGBColor.from_string(h.lstrip("#"))
def textbox(slide,x,y,w,h,text,size=20,color=TEXT,bold=False):
    shape=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=shape.text_frame;tf.word_wrap=True
    for i,line in enumerate(text.split("\n")):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line;p.font.name="Segoe UI";p.font.size=Pt(size);p.font.color.rgb=rgb(color);p.font.bold=bold
        p.space_after=Pt(12)
    return shape
def base(title,kicker,notes):
    slide=prs.slides.add_slide(prs.slide_layouts[6]);slide.background.fill.solid();slide.background.fill.fore_color.rgb=rgb(BG)
    bar=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(.55),Inches(.5),Inches(.08),Inches(.35));bar.fill.solid();bar.fill.fore_color.rgb=rgb(CYAN);bar.line.fill.background()
    textbox(slide,.8,.48,11.5,.35,kicker.upper(),11,CYAN,True)
    textbox(slide,.75,1.08,11.8,1.1,title,32,TEXT,True)
    textbox(slide,.75,7.03,10,.2,"PULSEROUTE / 28 OCTOBER 2026 / ENGINEERING CHECKPOINT",9,MUTED)
    textbox(slide,12,7.03,.6,.2,f"{len(prs.slides):02d}",9,MUTED)
    slide.notes_slide.notes_text_frame.text=notes
    return slide

def native_diagram(slide,image):
    from pptx.enum.shapes import MSO_CONNECTOR
    from pptx.oxml.xmlchemy import OxmlElement
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    ident=image.stem
    d=next(x for x in data["diagrams"] if x["id"]==ident)
    if ident.startswith("06"):
        xs=[1.5,4.7,7.9,11.1]
        for x,label in zip(xs,["Sender","Receiver","Disk","SQLite"]):
            textbox(slide,x-.65,2.15,1.6,.5,label,20,CYAN,True)
            line=slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x),Inches(2.8),Inches(x),Inches(6.4))
            line.line.color.rgb=rgb("#34465e")
        rows=[(0,1,"put: ID, name, size, digest"),(1,0,"ready"),(0,1,"chunks, then end"),(1,2,"verify; sync and rename"),(1,3,"record completion"),(1,0,"complete: ID, size, digest")]
        for i,(a,b,label) in enumerate(rows):
            y=3+i*.52
            line=slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(xs[a]),Inches(y),Inches(xs[b]),Inches(y))
            line.line.color.rgb=rgb(CYAN);line.line.width=Pt(1.5)
            end=OxmlElement("a:tailEnd");end.set("type","triangle");line.line._get_or_add_ln().append(end)
            textbox(slide,min(xs[a],xs[b])+.1,y-.34,abs(xs[b]-xs[a])-.1,.3,label,13,TEXT)
        textbox(slide,.85,6.5,11.5,.35,"Sender hashes first. Receiver checks content before committing. Full failure paths are in the report.",12,MUTED)
        return
    layouts={
     "01":[(1.7,2.9),(4.9,2.9),(8.1,2.9),(8.1,5.5),(4.9,5.5),(11.3,4.2)],
     "04":[(2,3.1),(6.6,3.1),(11.1,3.1),(3.6,5.55),(9.2,5.55)],
     "08":[(1.8,3),(6.6,3),(11.3,3),(4,5.4),(9.4,5.4)]
    }
    pos=layouts[ident[:2]]
    for a,b,label in d["edges"]:
        x1,y1=pos[a];x2,y2=pos[b];dx=x2-x1;dy=y2-y1
        scale=max(abs(dx)/1.21,abs(dy)/.62)
        line=slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1+dx/scale),Inches(y1+dy/scale),Inches(x2-dx/scale),Inches(y2-dy/scale))
        line.line.color.rgb=rgb(CYAN);line.line.width=Pt(1.6)
        end=OxmlElement("a:tailEnd");end.set("type","triangle");line.line._get_or_add_ln().append(end)
        textbox(slide,(x1+x2)/2-.9,(y1+y2)/2-(.98 if abs(dy)<.2 else .2),1.9,.55,label,11,MUTED)
    for (x,y),label in zip(pos,d["nodes"]):
        shape=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(x-1.2),Inches(y-.57),Inches(2.4),Inches(1.14))
        shape.fill.solid();shape.fill.fore_color.rgb=rgb(PANEL);shape.line.color.rgb=rgb("#34465e")
        tf=shape.text_frame;tf.word_wrap=True;tf.vertical_anchor=MSO_ANCHOR.MIDDLE
        tf.margin_left=Inches(.08);tf.margin_right=Inches(.08)
        for i,line in enumerate(label.replace("\\n","\n").split("\n")):
            para=tf.paragraphs[0] if i==0 else tf.add_paragraph();para.text=line;para.font.name="Segoe UI";para.font.size=Pt(13 if ident.startswith("04") else 17);para.font.color.rgb=rgb(TEXT);para.alignment=PP_ALIGN.CENTER
    textbox(slide,.85,6.65,11.5,.25,"Editable implementation diagram. Verification limitations are recorded separately.",11,MUTED)

def bullets(title,kicker,items,notes,image=None):
    slide=base(title,kicker,notes)
    if image and image.parent==FIG:
        native_diagram(slide,image)
        return slide
    width=5.4 if image else 11.4
    for i,item in enumerate(items):
        textbox(slide,.8,2.4+i*1.1,.25,.4,f"{i+1:02d}",11,CYAN,True)
        textbox(slide,1.2,2.32+i*1.1,width-.4,.92,item,20)
    if image and image.exists():
        from PIL import Image as PILImage
        w,h=PILImage.open(image).size
        maxw,maxh=5.5,4.7;rw=min(maxw,maxh*w/h);rh=rw*h/w
        slide.shapes.add_picture(str(image),Inches(7),Inches(2.15),width=Inches(rw),height=Inches(rh))
    return slide
slide=base("PulseRoute","Computer Networks laboratory","Introduce the project and state the evidence boundary immediately. The application is implemented in source and compiled, but current native execution is blocked by Smart App Control.")
textbox(slide,.8,2.65,10.8,1.8,"Move files.\nUnderstand the network.",43,TEXT,True)
textbox(slide,.85,5.35,11.5,.8,"Go + React + Wails + SQLite\nImplementation, verification and the remaining work",20,MUTED)
bullets("What the evidence supports","Scope",["Earlier real loopback messaging and file-identity tests passed.","Current frontend, SQL schema and build checks pass.","Native runtime and LAN performance remain BLOCKED / NOT VERIFIED."],"Do not describe compiled files as a fully tested application. Point to PROGRESS.md and the recorded evidence JSON.")
bullets("A workspace built around the flow","01 / Frontend",["Six screens: overview, messages, transfers, lab, history, settings.","Desktop controls call exported Go service methods.","Dark/light styling, focus states and explicit errors."],"Show the actual frontend screenshot. The bridge-unavailable banner is expected in browser-only verification and must not be hidden.",ROOT/"docs/evidence/frontend-overview.png")
bullets("Queue, observe, verify","01 / Frontend",["File picker and drop queue, then a destination endpoint.","Hashing, streaming, cancellation and final verification states.","Supported media preview; 16 MiB inline limit."],"Explain that the screenshot is a frontend rendering. Live desktop file picker, cancellation and media playback still need approved runtime validation.",ROOT/"docs/evidence/frontend-transfers.png")
bullets("Measure with clear boundaries","01 / Frontend",["TCP and UDP echo configuration share a laboratory view.","Charts render recorded samples only.","Export raw observations; empty charts remain empty."],"There is no synthetic benchmark curve. The absence of data is intentional.",ROOT/"docs/evidence/frontend-network-lab.png")
bullets("One service, separate engines","02 / Backend",["Wails coordinates state; transport packages own sockets.","Chat and file uploads use separate TCP connections.","SQLite persists outcomes; progress stays transient."],"Use the architecture diagram to explain separation of concerns.",FIG/"01-system-architecture.png")
bullets("Concurrency that can shut down","02 / Backend",["At most 64 active TCP handlers and bounded frames.","Per-connection serialized writes, deadlines and heartbeat checks.","Shutdown cancels tracked work before closing the database."],"Recent lifecycle fixes compile but still require race and runtime regression tests. Do not claim a leak-free soak test.")
bullets("Completion means verified content","02 / Backend",["Hash source; stream binary chunks to a temporary file.","Receiver checks declared size and SHA-256.","Sync, unique rename, history record, then final acknowledgment."],"A lost final acknowledgment can leave the sender uncertain even if the receiver has committed the file. Resume and crash reconciliation remain future work.",FIG/"06-transfer-sequence.png")
bullets("TCP and UDP expose different guarantees","03 / Algorithms and networking",["TCP is an ordered byte stream; PulseRoute adds framing.","UDP preserves datagram boundaries but may lose or reorder packets.","The OS owns TCP congestion control; no Reno/CUBIC implementation is claimed."],"Discuss transport versus application responsibilities. Video-file transfer is not live video streaming.")
bullets("Useful formulas, honest interpretation","03 / Algorithms and networking",["RTT = receive time - send time; loss uses unique replies.","RTT variation is smoothed with a 1/16 update factor.","Echo goodput is payload/time, not saturated link capacity."],"Sequential probes, timing windows and loopback environment limit interpretation. Refer to ALGORITHMS.md for exact formulas.")
bullets("Direct LAN versus chat hub","03 / Algorithms and networking",["Multiple clients can share a chat listener.","Direct files target the receiver's own TCP endpoint.","No NAT traversal; hub file forwarding is not yet implemented."],"Do not imply file relay through the hub. Two-device validation remains outstanding.",FIG/"08-network-topology.png")
bullets("Local, versioned persistence","04 / Database and results",["Five tables, indexes, parameters and a versioned migration.","Benchmark session and samples commit together.","Settings, known peers and recent results restore on startup."],"Python SQL schema tests passed. The Go driver and restart workflow remain unverified under the execution block.",FIG/"04-entity-relationship.png")
bullets("Observed correctness, not a speed claim","04 / Database and results",["8 clients in the earlier messaging integration test.","0 B, 1 B, 65,537 B and 16 MiB files verified twice each.","20/20 TCP and UDP loopback echoes in the earlier test."],"These figures come from successful test assertions. No measured RTT/goodput series was exported; do not turn these counts into performance claims.")
bullets("The execution blocker is identified","04 / Database and results",["Code Integrity event 3077, records 14737 and 14722.","Policy: VerifiedAndReputableDesktop / Smart App Control.","No security controls changed; code signing or an approved test environment is needed."],"The event XML confirms policy attribution. Signed build review is separate from source correctness. SAC has no individual-app exception.")
bullets("The next verification gate","Remaining work",["Run the current Go suite, race checks and desktop workflows.","Test two Windows devices and export repeatable raw benchmarks.","Complete hub file distribution, recovery and final review."],"GitHub still returns 404 and no remote commit is claimed. License choice remains pending. This deck is an interim engineering presentation.")
bullets("Live demonstration plan","Rehearsal",["Connect two peers and exchange text through a hub.","Transfer an image and video file; compare SHA-256.","Run both echo protocols, export data, then restart to inspect history."],"Use DEMO.md. If runtime approval is unavailable, present the source, screenshots and honest evidence boundary instead of pretending the live demo succeeded.")
prs.save(OUT/"PulseRoute-Presentation-DRAFT.pptx")
(OUT/"README.txt").write_text("PulseRoute academic artifacts - 28 October 2026\\nThese are progress/draft artifacts, not a claim of completed runtime verification.\\nTechnical report: DOCX and PDF. Presentation: PPTX with notes.\\nTen diagrams: docs/diagrams, in Mermaid, PNG and SVG.\\nSee docs/PROGRESS.md for PASS, FAIL, BLOCKED and NOT VERIFIED statuses.\\n",encoding="utf-8")
for svg in FIG.glob("*.svg"):
    svg.write_text("\n".join(line.rstrip() for line in svg.read_text(encoding="utf-8").splitlines())+"\n",encoding="utf-8")
manifest={"reportSections":len(data["sections"]),"slides":len(prs.slides),"diagrams":len(data["diagrams"]),"artifacts":[{"name":p.name,"bytes":p.stat().st_size} for p in OUT.iterdir() if p.is_file() and p.name!="artifact-manifest.json"]}
(OUT/"artifact-manifest.json").write_text(json.dumps(manifest,indent=2),encoding="utf-8")
print(json.dumps(manifest,indent=2))
