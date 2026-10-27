import json, pathlib, zipfile, hashlib
import pymupdf
from PIL import Image, ImageDraw
from docx import Document
from pptx import Presentation
root=pathlib.Path(__file__).resolve().parents[1]
out=root/"deliverables";evidence=root/"docs/evidence"
for suffix in ("docx","pptx"):
    p=next(out.glob("*."+suffix))
    with zipfile.ZipFile(p) as z:
        assert z.testzip() is None
doc=Document(out/"PulseRoute-Technical-Report-DRAFT.docx")
headings=[p.text for p in doc.paragraphs if p.style.name=="Heading 1"]
assert any("22. Appendices" in h for h in headings)
prs=Presentation(out/"PulseRoute-Presentation-DRAFT.pptx")
assert len(prs.slides)==16
assert all(s.notes_slide.notes_text_frame.text.strip() for s in prs.slides)
bounds=[]
for i,s in enumerate(prs.slides,1):
    for shape in s.shapes:
        if shape.left<0 or shape.top<0 or shape.left+shape.width>prs.slide_width+10000 or shape.top+shape.height>prs.slide_height+10000:
            bounds.append({"slide":i,"shape":shape.name})
assert not bounds,bounds
pdf=pymupdf.open(out/"PulseRoute-Technical-Report-DRAFT.pdf")
for page in pdf:
    for item in page.get_images():
        for rect in page.get_image_rects(item[0]):
            assert rect.x0>=0 and rect.y0>=0 and rect.x1<=page.rect.width and rect.y1<=page.rect.height, (page.number,rect)
alltext="\n".join(p.get_text() for p in pdf)
assert "22. Appendices" in alltext and "BLOCKED" in alltext and "14737" in alltext
thumbs=[]
for i,page in enumerate(pdf):
    pix=page.get_pixmap(matrix=pymupdf.Matrix(.48,.48),alpha=False)
    img=Image.frombytes("RGB",[pix.width,pix.height],pix.samples)
    canvas=Image.new("RGB",(310,440),"#e8edf2");canvas.paste(img,((310-img.width)//2,18))
    ImageDraw.Draw(canvas).text((15,420),f"Page {i+1}",fill="#182536")
    thumbs.append(canvas)
sheet=Image.new("RGB",(310*4,440*((len(thumbs)+3)//4)),"#d4dfe8")
for i,img in enumerate(thumbs):sheet.paste(img,((i%4)*310,(i//4)*440))
sheet.save(evidence/"report-contact-sheet.png")
native="NOT VERIFIED"
proof=evidence/"native-presentation-render.json"
if proof.exists():
    record=json.loads(proof.read_text(encoding="utf-8-sig"))
    actual=hashlib.sha256((out/"PulseRoute-Presentation-DRAFT.pptx").read_bytes()).hexdigest()
    if record.get("status")=="PASS" and record.get("sha256", "").lower()==actual and record.get("slides")==16:
        native="PASS: native PowerPoint rendered this exact presentation"
result={"status":"PASS","pdfPages":len(pdf),"reportSections":22,"slides":len(prs.slides),"speakerNotes":"present on every slide","officeArchives":"valid ZIP packages, opened by document parsers","slideBounds":"all shapes within slide","nativePowerPointRendering":native,"diagramFiles":len(list((root/"docs/diagrams").glob("*.mmd")))}
(evidence/"artifact-validation.json").write_text(json.dumps(result,indent=2),encoding="utf-8")
print(json.dumps(result,indent=2))
